"""Run with the bundled document Python. Preserve native Word mathematics."""
from pathlib import Path
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1];path=root/'B题鲁棒优化_论文交接.docx'
d=Document(path)
for s in d.sections:
 s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=s.bottom_margin=Cm(2.3);s.left_margin=s.right_margin=Cm(2.4)
 if not any(p.text for p in s.footer.paragraphs) and not s.footer._element.xpath('.//w:fldSimple'):
  p=s.footer.paragraphs[0];p.alignment=1
  field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
for name,size in [('Normal',11),('Title',20),('Subtitle',11),('Heading 1',15),('Heading 2',13),('Heading 3',11.5)]:
 st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'SimSun'if name=='Normal'else'Microsoft YaHei')
 st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.2
for p in d.paragraphs:
 if p.text=='Table of Contents':p.text='目录'
for text in d._element.iter(qn('w:t')):
 if text.text=='Table of Contents':text.text='目录'
if 'TOC Heading'in d.styles:
 d.styles['TOC Heading'].font.color.rgb=RGBColor(0,0,0)
for shape in d.inline_shapes:
 if shape.width>Cm(16.2):
  ratio=Cm(16.2)/shape.width;shape.width=int(shape.width*ratio);shape.height=int(shape.height*ratio)
for table in d.tables:
 for i,row in enumerate(table.rows):
  if i==0 and row._tr.get_or_add_trPr().find(qn('w:tblHeader'))is None:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
  for cell in row.cells:
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.05
    for run in p.runs:run.font.size=Pt(9.5)
d.save(path)
print('Word styles normalized with native math preserved')
