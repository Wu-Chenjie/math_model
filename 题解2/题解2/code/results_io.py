# -*- coding: utf-8 -*-
"""生成 result1/2/3/4-2/4-3.xlsx（严格遵循附件5模板的表头与行列结构）。
result1: 模板行数恰好(145行/7行)，直接填充。
result2/3/4: "计划购电量""调整购电量"按模板334天行填充；"充放电量""紧急购电量"
             模板行数不足，重建（表头文字与列序严格复制模板）。
"""
import sys
import numpy as np
import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

sys.path.insert(0, r'D:/shumo/shumoC/code')
from common import load_data, TBL2_RANGES
from strategy import merge_emerg_segments, hhmm, SPEC_DATES

TPL = r"C:/Users/y'h/Downloads/CUMCM2026Problems/C题/附件/附件5"
OUT = r'D:/shumo/shumoC/results'
D = load_data()
DT = 1 / 6
DATES = np.arange('2025-02-01', '2026-01-01', dtype='datetime64[D]')  # 334天(含12-31)
DAY_IDX = np.arange(31, 365)                                          # 对应数据d索引
BOLD = Font(bold=True)
CENTER = Alignment(horizontal='center')

def fill_day_ahead(ws_days):
    """计划/调整购电量sheet填充：模板已有表头+334日期行。"""
    pass  # 在具体函数中处理

def slot_labels():
    """规范144时段区间标签：'0:00-0:10' ... '23:50-0:00+1'。"""
    labs = []
    for t in range(144):
        labs.append(f'{hhmm(t)}-{hhmm(t + 1)}' if t < 143 else '23:50-0:00+1')
    return labs

def gen_result1():
    z = np.load(rf'{OUT}/p1.npz')
    p, c, d, S = z['p'], z['c'], z['d'], z['S']
    wb = openpyxl.load_workbook(rf'{TPL}/result1.xlsx')
    ws = wb['计划购电量']
    # 模板时段标签整体错位一格（缺'0:00-0:10'、末行多'0:00+1-0:10+1'），修正为规范序列
    for t, lab in enumerate(slot_labels()):
        ws.cell(row=2 + t, column=1, value=lab)
    for t in range(144):
        ws.cell(row=2 + t, column=2, value=round(float(p[t]), 2))
    ws = wb['充放电量']
    for i, (a, b) in enumerate(TBL2_RANGES):
        ws.cell(row=2 + i, column=2, value=round(float(c[a:b].sum()), 2))
        ws.cell(row=2 + i, column=3, value=round(float(d[a:b].sum()), 2))
    ws.cell(row=2, column=5, value=6000.0)   # 0:00 储电量
    ws.cell(row=3, column=5, value=round(float(S[-1]), 2))  # 24:00
    wb.save(rf'{OUT}/result1.xlsx')
    print('result1.xlsx 完成')

def _write_planning_sheet(ws, p_days, price_mode):
    """向 计划/调整购电量 sheet 的334天行写入144列购电量与全天统计。
    模板时段列表头错位一格（同result1），先修正为规范序列。"""
    for t, lab in enumerate(slot_labels()):
        ws.cell(row=1, column=2 + t, value=lab)
    price1, price4 = D['price1'], D['price4']
    for i in range(334):
        di = DAY_IDX[i]
        row = 2 + i
        for t in range(144):
            ws.cell(row=row, column=2 + t, value=round(float(p_days[i, t]), 2))
        price = price1 if price_mode == 'fixed' else price4[di]
        ws.cell(row=row, column=146, value=round(float(p_days[i].sum()), 2))     # 全天购电量
        ws.cell(row=row, column=147, value=round(float(price @ p_days[i]), 2))   # 全天购电费

def _rebuild_charge_sheet(wb, c_days, d_days, sheet_name='充放电量'):
    """重建充放电量sheet：表头+334天×6行。"""
    hdr = [c.value for c in wb[sheet_name][1]]
    del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    ws.append(hdr)
    for cell in ws[1]:
        cell.font = BOLD; cell.alignment = CENTER
    for i in range(334):
        date = str(DATES[i])
        base = ws.max_row + 1
        for k, (a, b) in enumerate(TBL2_RANGES):
            r = base + k
            ws.cell(row=r, column=1, value=date if k == 0 else None)
            ws.cell(row=r, column=2, value=['0:00-4:00', '4:00-8:00', '8:00-12:00',
                                            '12:00-16:00', '16:00-20:00', '20:00-24:00'][k])
            ws.cell(row=r, column=3, value=round(float(c_days[i, a:b].sum()), 2))
            ws.cell(row=r, column=4, value=round(float(d_days[i, a:b].sum()), 2))
        S = 6000.0 + np.cumsum(0.9 * c_days[i] - d_days[i])
        ws.cell(row=base, column=5, value='00:00')
        ws.cell(row=base, column=6, value=6000.0)
        ws.cell(row=base + 1, column=5, value='24:00')
        ws.cell(row=base + 1, column=6, value=round(float(S[-1]), 2))

def _rebuild_emerg_sheet(wb, e_days, sheet_name='紧急购电量'):
    """重建紧急购电量sheet：日期+连续时段段+电量。"""
    hdr = [c.value for c in wb[sheet_name][1]]
    del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    ws.append(hdr)
    for cell in ws[1]:
        cell.font = BOLD; cell.alignment = CENTER
    n_segs = 0
    for i in range(334):
        segs = merge_emerg_segments(e_days[i])
        if not segs:
            ws.cell(row=ws.max_row + 1, column=1, value=str(DATES[i]))
        else:
            for (t0, t1, q) in segs:
                r = ws.max_row + 1
                ws.cell(row=r, column=1, value=str(DATES[i]) if r == ws.max_row or
                        ws.cell(row=r - 1, column=1).value != str(DATES[i]) else None)
                ws.cell(row=r, column=2, value=f'{hhmm(t0)}-{hhmm(t1)}')
                ws.cell(row=r, column=3, value=round(float(q), 2))
                n_segs += 1
    return n_segs

def gen_result2(npz, fname, price_mode='fixed'):
    z = np.load(npz)
    p, c, d, e = z['p'][31:], z['c'][31:], z['dd'][31:], z['e'][31:]
    tpl = 'result2.xlsx' if price_mode == 'fixed' else 'result4-2.xlsx'
    wb = openpyxl.load_workbook(rf'{TPL}/{tpl}')
    _write_planning_sheet(wb['计划购电量'], p, price_mode)
    _rebuild_charge_sheet(wb, c, d)
    n = _rebuild_emerg_sheet(wb, e)
    wb.save(rf'{OUT}/{fname}')
    print(f'{fname} 完成（紧急购电段数={n}）')

def gen_result3(npz, fname, price_mode='fixed'):
    z = np.load(npz)
    p, c, d, e = z['p'][31:], z['c'][31:], z['dd'][31:], z['e'][31:]
    plan_p = z['plan_p'][31:]
    tpl = 'result3.xlsx' if price_mode == 'fixed' else 'result4-3.xlsx'
    wb = openpyxl.load_workbook(rf'{TPL}/{tpl}')
    _write_planning_sheet(wb['计划购电量'], plan_p, price_mode)   # 计划=0:00计划
    _write_planning_sheet(wb['调整购电量'], p, price_mode)        # 调整=最终执行
    _rebuild_charge_sheet(wb, c, d)
    n = _rebuild_emerg_sheet(wb, e)
    wb.save(rf'{OUT}/{fname}')
    print(f'{fname} 完成（紧急购电段数={n}）')

if __name__ == '__main__':
    gen_result1()
    gen_result2(rf'{OUT}/p2.npz', 'result2.xlsx', 'fixed')
    gen_result3(rf'{OUT}/p3.npz', 'result3.xlsx', 'fixed')
    gen_result2(rf'{OUT}/p4_2.npz', 'result4-2.xlsx', 'realtime')
    gen_result3(rf'{OUT}/p4_3.npz', 'result4-3.xlsx', 'realtime')
